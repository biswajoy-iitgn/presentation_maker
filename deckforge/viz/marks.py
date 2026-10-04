"""Drawing primitives for shape-built exhibits. Thin marks, no borders, text in ink tokens.

Colour and size defaults resolve at call time, so they follow the active style family.
"""

from __future__ import annotations

import functools

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from deckforge.render import metrics
from deckforge.render.canvas import Canvas, rgb
from deckforge.viz import style as S

GAP = Inches(0.028)      # surface gap between adjacent fills (about 2 px at slide scale)


def bar(c: Canvas, x, y, w, h, color, name=None, register=True, shape=MSO_SHAPE.RECTANGLE):
    """Filled rectangle without border. Negative heights are normalised.
    Registered as a fill so the legibility lint checks text drawn on it."""
    if h < 0:
        y, h = y + h, -h
    if w < 0:
        x, w = x + w, -w
    return c.rect(int(x), int(y), max(1, int(w)), max(1, int(h)), color, name=name, register=register, shape=shape)


def label(c: Canvas, x, y, w, h, s, *, size=None, bold=False, color=None, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, name=None, font=None):
    return c.text(int(x), int(y), int(w), int(h), s, size=size or S.TYPE.label, bold=bold, color=color or S.INK,
                  align=align, anchor=anchor, name=name, font=font or S.FONT)


def text_w(s: str, size=None, bold=False, font=None) -> int:
    return int(metrics.text_width_pt(s, font or S.FONT, size or S.TYPE.label, bold) * 12700)


def centered_label(c: Canvas, cx, y, s, *, w=Inches(1.4), h=Inches(0.28), knockout=False, **kw):
    """knockout: white box sized to the text, so lines running behind the label do not cut through it."""
    if knockout and isinstance(s, str):
        tw = text_w(s, kw.get("size"), kw.get("bold", False), kw.get("font")) + Inches(0.08)
        c.rect(int(cx - tw / 2), int(y + Inches(0.02)), tw, int(h - Inches(0.04)), "#FFFFFF", register=False)
    return label(c, cx - w // 2, y, w, h, s, align=PP_ALIGN.CENTER, **kw)


def hline(c: Canvas, x1, x2, y, color=None, width=0.75):
    return c.line(int(x1), int(y), int(x2), int(y), color or S.RULE, width)


def vline(c: Canvas, x, y1, y2, color=None, width=0.75):
    return c.line(int(x), int(y1), int(x), int(y2), color or S.RULE, width)


def dot(c: Canvas, cx, cy, d, color, ring=None, name=None, register=False):
    return c.rect(int(cx - d / 2), int(cy - d / 2), int(d), int(d), color, shape=MSO_SHAPE.OVAL, name=name,
                  register=register, line=(ring, 1.5) if ring else None)


def _shape_text(sp, text, size, color, bold=True):
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    tf.text = text
    run = tf.paragraphs[0].runs[0]
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, S.FONT
    run.font.color.rgb = rgb(color)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER


def marker(c: Canvas, cx, cy, n: int, d=Inches(0.29), color=None):
    """Numbered callout marker linking a chart element to commentary."""
    sp = dot(c, cx, cy, d, color or S.INK, ring="#FFFFFF", name=f"callout_{n}")
    _shape_text(sp, str(n), 11, "#FFFFFF")
    return sp


def badge(c: Canvas, x, y, text, *, fill=None, color="#FFFFFF", size=None, h=Inches(0.34), pad=Inches(0.14),
          name="badge") -> int:
    """Rounded pill carrying a short emphasised statement, e.g. '-390 bps'. Returns its width."""
    size = size or S.TYPE.value
    w = text_w(text, size, True) + 2 * pad
    sp = c.rect(int(x), int(y), int(w), int(h), fill or S.ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=name,
                register=False)
    sp.adjustments[0] = 0.5
    _shape_text(sp, text, size, color)
    return w


@functools.lru_cache(maxsize=256)
def _icon_png(name: str, color: str, px: int, stroke: float):
    from deckforge.assets.sources.icons import icon_png
    img, _ = icon_png(name, color, px, stroke_width=stroke)
    return img


def icon(c: Canvas, name: str, x, y, size, color=None, stroke: float = 1.75):
    """Line icon from the open icon set, recoloured to the family, placed as a transparent PNG."""
    img = _icon_png(name, color or S.ACCENT, 256, stroke)
    return c.picture(img, int(x), int(y), int(size), int(size), name=f"icon_{name}", fmt="PNG")


def icon_disc(c: Canvas, name: str, cx, cy, d, *, fill=None, color=None):
    """Icon on a filled circle: the standard consulting category marker.

    The glyph colour follows the fill: white on dark or saturated fills, ink on light neutrals, so a
    de-emphasised marker still reads (a white glyph on light grey is under 2:1 and disappears).
    """
    fill = fill or S.ACCENT
    dot(c, cx, cy, d, fill)
    if color is None:
        from deckforge.assets.treatment import luminance
        color = S.INK if luminance(fill) > 0.45 else "#FFFFFF"
    s = d * 0.56
    return icon(c, name, cx - s / 2, cy - s / 2, s, color)


def kpi(c: Canvas, x, y, w, value: str, caption: str, *, color=None, caption_color=None, size=None,
        name="kpi") -> int:
    """Big number over a short caption. Returns the height used."""
    size = size or S.TYPE.kpi
    vh = int(size * 1.15 * 12700)
    label(c, x, y, w, vh, value, size=size, bold=True, color=color or S.ACCENT, name=f"{name}_value")
    ch = metrics.text_height(caption, S.FONT, S.TYPE.label, False, int(w))
    label(c, x, y + vh + Inches(0.02), w, ch, caption, size=S.TYPE.label, color=caption_color or S.TEXT2,
          name=f"{name}_caption")
    return vh + Inches(0.02) + ch


def bracket_h(c: Canvas, x1, x2, y, label_text, *, up=True, color=None, tick=Inches(0.08), size=None, bold=True):
    """Horizontal bracket spanning x1..x2 with a centred label (above if up)."""
    color = color or S.INK
    size = size or S.TYPE.annotation
    t = -tick if up else tick
    c.line(int(x1), int(y + t), int(x1), int(y), color, 1.0)
    c.line(int(x1), int(y), int(x2), int(y), color, 1.0)
    c.line(int(x2), int(y), int(x2), int(y + t), color, 1.0)
    ly = y - Inches(0.32) if up else y + Inches(0.06)
    tw = text_w(label_text, size, bold) + Inches(0.15)
    return centered_label(c, (x1 + x2) // 2, ly, label_text, w=max(tw, int(x2 - x1)), size=size, bold=bold,
                          color=color)


def arrow(c: Canvas, x1, y1, x2, y2, color=None, width=1.0):
    return c.line(int(x1), int(y1), int(x2), int(y2), color or S.INK, width, arrow_end=True)


def ring(c: Canvas, cx, cy, d, share: float, *, color=None, track=None, thickness=0.22, name="ring"):
    """Donut ring showing a share (0..1) as native shapes, starting at 12 o'clock."""
    color, track = color or S.ACCENT, track or S.TRACK
    x, y = int(cx - d / 2), int(cy - d / 2)
    full = c.rect(x, y, int(d), int(d), track, shape=MSO_SHAPE.DONUT, register=False, name=f"{name}_track")
    full.adjustments[0] = thickness
    if share > 0:
        # python-pptx exposes OOXML angles (60000ths of a degree) divided by 100000, so degrees x 0.6
        arc = c.rect(x, y, int(d), int(d), color, shape=MSO_SHAPE.BLOCK_ARC, register=False, name=f"{name}_value")
        arc.adjustments[0] = 270.0 * 0.6
        arc.adjustments[1] = ((270.0 + 360.0 * min(share, 0.9999)) % 360) * 0.6
        arc.adjustments[2] = thickness
    return full
