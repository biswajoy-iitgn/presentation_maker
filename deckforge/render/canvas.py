"""Deck and slide canvas: every placement goes through here so QA knows exactly what is where."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from deckforge.render import metrics
from deckforge.tokens import Theme

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#"))


Box = tuple[int, int, int, int]   # x, y, w, h in EMU


@dataclass
class Placed:
    """A placement record used by the geometry and legibility lint."""
    kind: str                    # text | sticker | image | fill
    name: str
    box: Box
    color: str | None = None     # text colour or fill colour
    size: float | None = None
    bold: bool = False
    image: Image.Image | None = None   # effective background (scrims composited in)
    shape: object | None = None


@dataclass
class Canvas:
    deck: "Deck"
    slide: object
    index: int
    placed: list[Placed] = field(default_factory=list)

    @property
    def theme(self) -> Theme:
        return self.deck.theme

    # ------------------------------------------------------------------ text
    def text(self, x, y, w, h, content, *, font=None, size=12, bold=False, color=None,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None, leading=None):
        """content: str, or list of paragraphs, each str or list of (text, bold[, color]) runs."""
        font = font or self.theme.body_font
        color = color or self.theme.ink
        paras = content if isinstance(content, list) else [content]
        plain = "\n".join(p if isinstance(p, str) else "".join(r[0] for r in p) for p in paras)
        any_bold = bold or any(not isinstance(p, str) and any(r[1] for r in p) for p in paras)
        used = metrics.text_height(plain, font, size, any_bold, w)
        tb = self.slide.shapes.add_textbox(x, y, w, h)
        self.placed.append(Placed("text", name or plain[:40], (x, y, w, used if anchor == MSO_ANCHOR.TOP else h),
                                  color, size, any_bold, shape=tb))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if leading:
                p.line_spacing = leading
            runs = [(para, bold)] if isinstance(para, str) else para
            for run_spec in runs:
                seg, seg_bold = run_spec[0], run_spec[1]
                seg_color = run_spec[2] if len(run_spec) > 2 else color
                r = p.add_run()
                r.text = seg
                r.font.name = font
                r.font.size = Pt(size)
                r.font.bold = seg_bold
                r.font.color.rgb = rgb(seg_color)
        if name:
            tb.name = name
        return tb

    # ------------------------------------------------------------------ shapes
    def rect(self, x, y, w, h, fill, *, line=None, shape=MSO_SHAPE.RECTANGLE, name=None, alpha=None,
             register=True):
        sp = self.slide.shapes.add_shape(shape, x, y, w, h)
        sp.shadow.inherit = False
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = rgb(fill)
            if alpha is not None:
                clr = sp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = rgb(line[0])
            sp.line.width = Pt(line[1])
        if name:
            sp.name = name
        if register and fill is not None and (alpha is None or alpha >= 0.95):
            self.placed.append(Placed("fill", name or "rect", (x, y, w, h), fill))
        return sp

    def sticker(self, label, x, y, w, h, color):
        sp = self.rect(x, y, w, h, self.theme.paper, line=(color, 1.25), name="DF_ILLUSTRATIVE", register=False)
        tf = sp.text_frame
        tf.text = label
        run = tf.paragraphs[0].runs[0]
        run.font.size, run.font.bold, run.font.name = Pt(11), True, self.theme.body_font
        run.font.color.rgb = rgb(color)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        self.placed.append(Placed("sticker", "DF_ILLUSTRATIVE", (x, y, w, h)))
        return sp

    def line(self, x1, y1, x2, y2, color, width=0.75, arrow_end=False, dash=False):
        c = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(width)
        ln = c.line._get_or_add_ln()
        if dash:
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
        if arrow_end:
            tail = etree.SubElement(ln, qn("a:tailEnd"))
            tail.set("type", "triangle")
            tail.set("w", "med")
            tail.set("len", "med")
        return c

    def freeform(self, contours: list[list[tuple[int, int]]], fill, *, line=None, name=None):
        """Native editable polygon(s), e.g. a country with islands."""
        first = contours[0]
        fb = self.slide.shapes.build_freeform(first[0][0], first[0][1], scale=1.0)
        fb.add_line_segments(first[1:], close=True)
        for c in contours[1:]:
            fb.move_to(c[0][0], c[0][1])
            fb.add_line_segments(c[1:], close=True)
        sp = fb.convert_to_shape()
        sp.shadow.inherit = False
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
        if line:
            sp.line.color.rgb = rgb(line[0])
            sp.line.width = Pt(line[1])
        else:
            sp.line.fill.background()
        if name:
            sp.name = name
        return sp

    # ------------------------------------------------------------------ images
    def picture(self, img: Image.Image | str | Path, x, y, w, h, *, name=None, fmt="JPEG"):
        """Place an image. The effective image is kept for legibility checks under text."""
        pil = Image.open(img).convert("RGB") if isinstance(img, (str, Path)) else img
        buf = io.BytesIO()
        if fmt == "PNG" or pil.mode == "RGBA":
            pil.save(buf, "PNG", optimize=True)
        else:
            pil.convert("RGB").save(buf, "JPEG", quality=88, optimize=True)
        buf.seek(0)
        pic = self.slide.shapes.add_picture(buf, x, y, w, h)
        if name:
            pic.name = name
        if pil.mode != "RGBA":
            self.placed.append(Placed("image", name or "image", (x, y, w, h), image=pil.convert("RGB").copy()))
        return pic

    def scrim(self, x, y, w, h, color, alpha_from, alpha_to, angle_deg=0, name="DF_scrim"):
        """Native gradient overlay with transparency. Also composited into images below for QA."""
        sp = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        sp.shadow.inherit = False
        sp.line.fill.background()
        sp_pr = sp._element.spPr
        for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
            for el in sp_pr.findall(qn(tag)):
                sp_pr.remove(el)
        grad = etree.Element(qn("a:gradFill"), rotWithShape="1")
        gs_lst = etree.SubElement(grad, qn("a:gsLst"))
        for pos, a in ((0, alpha_from), (100000, alpha_to)):
            gs = etree.SubElement(gs_lst, qn("a:gs"), pos=str(pos))
            clr = etree.SubElement(gs, qn("a:srgbClr"), val=color.lstrip("#"))
            etree.SubElement(clr, qn("a:alpha"), val=str(int(a * 100000)))
        etree.SubElement(grad, qn("a:lin"), ang=str(int(angle_deg * 60000)), scaled="0")
        sp_pr.find(qn("a:prstGeom")).addnext(grad)
        sp.name = name
        self._composite_scrim((x, y, w, h), color, alpha_from, alpha_to, angle_deg)
        return sp

    def backlight(self, target: Placed, color: str, alpha: float, pad=Inches(0.35)):
        """Feathered, semi-transparent panel inserted directly behind a text shape (legibility repair)."""
        x, y, w, h = target.box
        box = (x - pad, y - pad, w + 2 * pad, h + 2 * pad)
        sp = self.rect(*box, color, alpha=alpha, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name="DF_backlight",
                       register=False)
        effects = etree.SubElement(sp._element.spPr, qn("a:effectLst"))
        etree.SubElement(effects, qn("a:softEdge"), rad=str(int(pad * 0.8)))
        target.shape._element.addprevious(sp._element)
        self._composite_scrim(box, color, alpha, alpha, 0)
        return sp

    def _composite_scrim(self, box: Box, color, a0, a1, angle_deg):
        from deckforge.assets.treatment import composite_gradient
        for p in self.placed:
            if p.kind == "image" and _intersects(p.box, box):
                p.image = composite_gradient(p.image, p.box, box, color, a0, a1, angle_deg)


def _intersects(a: Box, b: Box) -> bool:
    return a[0] < b[0] + b[2] and b[0] < a[0] + a[2] and a[1] < b[1] + b[3] and b[1] < a[1] + a[3]


class Deck:
    def __init__(self, theme: Theme, width=SLIDE_W, height=SLIDE_H):
        self.theme = theme
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = width, height
        self.width, self.height = width, height
        self.slides: list[Canvas] = []

    def new_slide(self) -> Canvas:
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        c = Canvas(self, s, len(self.slides) + 1)
        self.slides.append(c)
        return c

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        return path
