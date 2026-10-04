"""Look-and-feel gate: does the deck look and read like a top-firm deck, not like generated slides?

Thresholds come from the reviewed corpus decks (docs/corpus_notes):
- Accenture investor deck: imagery on all 17 slides (backgrounds, 3D graphics, map, icons).
- BCG executive perspectives: photo header band on every content slide, icons on most.
- McKinsey survey deck: icons or flags and a dark big-number sidebar on most slides, no photos.
- Bain PE report: photo cover and summary, icons on timelines, charts elsewhere.
So a premium consulting deck carries imagery or graphics on most slides, varies its layouts, gives every
content slide a focal element, keeps text in check, and uses chart forms its audience reads without decoding.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckforge.render.canvas import Deck
from deckforge.viz import style as S
from deckforge.viz.select import FAMILIARITY


@dataclass
class Rule:
    code: str
    passed: bool
    detail: str
    hard: bool = True


@dataclass
class LookFeel:
    rules: list[Rule] = field(default_factory=list)
    per_slide: list[dict] = field(default_factory=list)

    @property
    def passed(self) -> bool:
        return all(r.passed for r in self.rules if r.hard)

    def failures(self) -> list[Rule]:
        return [r for r in self.rules if not r.passed]

    def markdown(self) -> str:
        out = ["| Check | Result | Detail |", "|---|---|---|"]
        for r in self.rules:
            res = "pass" if r.passed else ("FAIL" if r.hard else "warn")
            out.append(f"| {r.code} | {res} | {r.detail} |")
        out += ["", "| Slide | Layout | Imagery | Graphics | Focal element | Words |", "|---|---|---|---|---|---|"]
        for s in self.per_slide:
            out.append(f"| {s['n']} | {s['layout']} | {s['images']} | {s['graphics']} | "
                       f"{'yes' if s['focal'] else 'no'} | {s['words']} |")
        return "\n".join(out)


GRAPHIC_PREFIXES = ("icon_", "map_", "site_", "bubble_", "callout_", "badge", "ring", "kpi_sidebar", "cover_art",
                    "agenda_panel", "DF_band")


def _shape_stats(slide) -> dict:
    images = graphics = words = 0
    accent_fills = 0
    min_font = 99.0
    for sh in slide.shapes:
        name = sh.name or ""
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            if name.startswith("icon_"):
                graphics += 1
            else:
                images += 1
        elif name.startswith(GRAPHIC_PREFIXES):
            graphics += 1
        if sh.has_text_frame and sh.text_frame.text.strip():
            words += len(re.findall(r"[A-Za-z][A-Za-z'-]*", sh.text_frame.text))     # numbers are not reading load
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and not name.startswith(("DF_foot", "DF_page", "DF_tracker", "cover_note")):
                        min_font = min(min_font, r.font.size.pt)
        try:
            if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == S.ACCENT.lstrip("#").upper():
                accent_fills += 1
        except (AttributeError, TypeError, ValueError):
            pass
    return {"images": images, "graphics": graphics, "words": words, "accent": accent_fills, "min_font": min_font}


def check(deck: Deck, decisions, *, audience: str = "board", min_visual_share: float = 0.7,
          min_layouts: int = 6, max_repeat: int = 2, max_words: int = 150) -> LookFeel:
    lf = LookFeel()
    layouts = [d.layout for d in decisions]
    visual = 0
    no_focal, wordy, small = [], [], []
    for i, (c, d) in enumerate(zip(deck.slides, decisions), 1):
        st = _shape_stats(c.slide)
        names = " ".join(sh.name or "" for sh in c.slide.shapes)
        focal = st["accent"] > 0 or "kpi" in names or "badge" in names or "callout_" in names
        is_visual = st["images"] + st["graphics"] > 0
        visual += is_visual
        content = d.archetype in ("exhibit", "exec_summary", "decisions")
        if content and not focal:
            no_focal.append(i)
        limit = max_words + (60 if d.archetype == "exec_summary" else 0)
        if st["words"] > limit:
            wordy.append(f"{i} ({st['words']})")
        if st["min_font"] < 9.5:
            small.append(f"{i} ({st['min_font']:g} pt)")
        lf.per_slide.append({"n": i, "layout": d.layout, "images": st["images"], "graphics": st["graphics"],
                             "focal": focal, "words": st["words"]})

    n = len(deck.slides)
    share = visual / n if n else 0
    lf.rules.append(Rule("imagery_or_graphics", share >= min_visual_share,
                         f"{visual} of {n} slides carry imagery or graphics ({share:.0%}, need {min_visual_share:.0%})"))
    cover_ok = lf.per_slide and lf.per_slide[0]["images"] > 0
    lf.rules.append(Rule("cover_imagery", bool(cover_ok), "cover carries imagery" if cover_ok else "cover has no imagery"))
    distinct = len(set(layouts))
    lf.rules.append(Rule("layout_variety", distinct >= min(min_layouts, n),
                         f"{distinct} distinct layouts (need {min(min_layouts, n)})"))
    run, worst, at = 1, 1, None
    for a, b in zip(layouts, layouts[1:]):
        run = run + 1 if a == b else 1
        if run > worst:
            worst, at = run, b
    lf.rules.append(Rule("layout_repeat", worst <= max_repeat,
                         f"longest run of one layout: {worst}" + (f" ({at})" if at and worst > 1 else "")))
    lf.rules.append(Rule("focal_element", not no_focal,
                         "every content slide has a focal element" if not no_focal else f"no focal element on slides {no_focal}"))
    lf.rules.append(Rule("text_density", not wordy, "text within limits" if not wordy else f"too many words: {wordy}"))
    lf.rules.append(Rule("min_font_size", not small, "all text 9.5 pt or larger" if not small else f"small text: {small}"))

    unfamiliar = [f"{d.slide} {d.chosen}" for d in decisions if d.chosen and FAMILIARITY.get(d.chosen, 0) < -0.02]
    lf.rules.append(Rule("chart_familiarity", not unfamiliar or audience not in ("board", "executive"),
                         "forms suit the audience" if not unfamiliar else f"specialist forms for a {audience}: {unfamiliar}",
                         hard=audience in ("board", "executive")))
    return lf
